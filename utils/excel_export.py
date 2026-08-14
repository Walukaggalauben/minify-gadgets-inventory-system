from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font


class ExcelExporter:

    @staticmethod
    def sales_report(sales):

        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        ws.append(["MINIFY GADGETS"])
        ws.append(["Sales Report"])
        ws.append([])

        headers = [
            "Invoice",
            "Customer",
            "Date",
            "Total",
            "Profit"
        ]

        ws.append(headers)

        for cell in ws[4]:
            cell.font = Font(bold=True)

        for sale in sales:

            ws.append([
                sale.invoice_number,
                sale.customer_name,
                sale.sale_date.strftime("%d/%m/%Y"),
                float(sale.total_amount),
                float(sale.profit)
            ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output
    
        # ==========================================================
    # PURCHASE REPORT
    # ==========================================================

    @staticmethod
    def purchase_report(purchases):

        wb = Workbook()
        ws = wb.active
        ws.title = "Purchase Report"

        ws.append(["MINIFY GADGETS"])
        ws.append(["Purchase Report"])
        ws.append([])

        headers = [
            "Purchase No",
            "Supplier",
            "Date",
            "Status",
            "Total"
        ]

        ws.append(headers)

        for cell in ws[4]:
            cell.font = Font(bold=True)

        for purchase in purchases:

            ws.append([
                purchase.purchase_number,
                purchase.supplier.name if purchase.supplier else "",
                purchase.purchase_date.strftime("%d/%m/%Y"),
                purchase.status,
                float(purchase.total_amount)
            ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output
    
           # ==========================================================
    # INVENTORY REPORT
    # ==========================================================

    @staticmethod
    def inventory_report(variants):

        wb = Workbook()
        ws = wb.active
        ws.title = "Inventory Report"

        # ------------------------------------------------------
        # TITLE
        # ------------------------------------------------------

        ws.append(["MINIFY GADGETS"])
        ws.append(["Inventory Report"])
        ws.append([])

        # ------------------------------------------------------
        # HEADERS
        # ------------------------------------------------------

        headers = [
            "Product",
            "Variant",
            "SKU",
            "Colour",
            "Storage",
            "Buying Price",
            "Selling Price",
            "Quantity",
            "Cost Value",
            "Selling Value",
            "Expected Profit",
        ]

        ws.append(headers)

        for cell in ws[4]:
            cell.font = Font(bold=True)

        # ------------------------------------------------------
        # INVENTORY DATA
        # ------------------------------------------------------

        for item in variants:

            quantity = item.quantity or 0
            buying_price = float(item.buying_price or 0)
            selling_price = float(item.selling_price or 0)

            cost_value = buying_price * quantity
            selling_value = selling_price * quantity
            expected_profit = selling_value - cost_value

            variant_parts = []

            if item.storage:
                variant_parts.append(str(item.storage))

            if item.ram:
                variant_parts.append(str(item.ram))

            if item.colour:
                variant_parts.append(str(item.colour))

            variant = " / ".join(variant_parts)

            ws.append([
                item.product.name if item.product else "",
                variant,
                item.sku or "",
                item.colour or "",
                item.storage or "",
                buying_price,
                selling_price,
                quantity,
                cost_value,
                selling_value,
                expected_profit,
            ])

        # ------------------------------------------------------
        # NUMBER FORMATTING
        # ------------------------------------------------------

        for row in ws.iter_rows(min_row=5):

            for cell in row:

                if cell.column in [6, 7, 9, 10, 11]:
                    cell.number_format = '#,##0.00'

                elif cell.column == 8:
                    cell.number_format = '#,##0'

        # ------------------------------------------------------
        # COLUMN WIDTHS
        # ------------------------------------------------------

        widths = {
            "A": 28,
            "B": 22,
            "C": 18,
            "D": 15,
            "E": 15,
            "F": 16,
            "G": 16,
            "H": 12,
            "I": 18,
            "J": 18,
            "K": 18,
        }

        for column, width in widths.items():
            ws.column_dimensions[column].width = width

        # ------------------------------------------------------
        # FREEZE HEADER
        # ------------------------------------------------------

        ws.freeze_panes = "A5"

        # ------------------------------------------------------
        # OUTPUT
        # ------------------------------------------------------

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output
    
        # ==========================================================
    # PROFIT REPORT
    # ==========================================================

    @staticmethod
    def profit_report(sales):

        wb = Workbook()
        ws = wb.active
        ws.title = "Profit Report"

        ws.append(["MINIFY GADGETS"])
        ws.append(["Profit Report"])
        ws.append([])

        ws.append([
            "Invoice",
            "Customer",
            "Date",
            "Sales",
            "Profit"
        ])

        for cell in ws[4]:
            cell.font = Font(bold=True)

        for sale in sales:
            ws.append([
                sale.invoice_number,
                sale.customer_name,
                sale.sale_date.strftime("%d/%m/%Y"),
                float(sale.total_amount),
                float(sale.profit)
            ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output
    
        # ==========================================================
    # IMEI REPORT
    # ==========================================================

    @staticmethod
    def imei_report(imeis):

        wb = Workbook()
        ws = wb.active
        ws.title = "IMEI Report"

        ws.append(["MINIFY GADGETS"])
        ws.append(["IMEI Report"])
        ws.append([])

        ws.append([
            "IMEI",
            "Product",
            "SKU",
            "Storage",
            "Colour",
            "Status"
        ])

        for cell in ws[4]:
            cell.font = Font(bold=True)

        for item in imeis:

            variant = item.product_variant

            ws.append([
                item.imei,
                variant.product.name if variant else "",
                variant.sku if variant else "",
                variant.storage if variant else "",
                variant.colour if variant else "",
                item.status
            ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output
    
        # ==========================================================
    # LOW STOCK REPORT
    # ==========================================================

    @staticmethod
    def low_stock_report(products):

        wb = Workbook()
        ws = wb.active
        ws.title = "Low Stock Report"

        ws.append(["MINIFY GADGETS"])
        ws.append(["Low Stock Report"])
        ws.append([])

        ws.append([
            "Product",
            "SKU",
            "Colour",
            "Storage",
            "Current Stock",
            "Minimum Stock"
        ])

        for cell in ws[4]:
            cell.font = Font(bold=True)

        for item in products:
            ws.append([
                item.product.name,
                item.sku,
                item.colour,
                item.storage,
                item.quantity,
                item.minimum_stock
            ])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output